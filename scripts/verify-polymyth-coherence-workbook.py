#!/usr/bin/env python3
"""Verify the byte-exact, blank Polymyth Coherence V5.1.2 release."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit(
        "POLYMYTH COHERENCE WORKBOOK FAILED — openpyxl is required"
    ) from exc


ROOT = Path(__file__).resolve().parent.parent
SITE_ONLY = "--site-only" in sys.argv[1:]
UNKNOWN_ARGUMENTS = [argument for argument in sys.argv[1:] if argument != "--site-only"]
if UNKNOWN_ARGUMENTS:
    raise SystemExit(
        "POLYMYTH COHERENCE WORKBOOK FAILED — unsupported argument(s): "
        + ", ".join(UNKNOWN_ARGUMENTS)
    )

PACKAGE_ROOT = ROOT.parent
COHERENCE_DIR = ROOT / "polymyth" / "coherence"
PUBLIC_COHERENCE_DIR = ROOT / "public" / "polymyth" / "coherence"
EDITABLE_DIR = PACKAGE_ROOT / "EDITABLE_MASTERS" / "07_POLYMYTH_COHERENCE"

INSTRUMENT = COHERENCE_DIR / "Polymyth_Coherence_Assessment_Instrument_V5.1.2.xlsx"
PUBLIC_INSTRUMENT = PUBLIC_COHERENCE_DIR / "Polymyth_Coherence_Assessment_Instrument_V5.1.2.xlsx"
PROTOCOL = COHERENCE_DIR / "Polymyth_Coherence_AI_Application_Protocol_V5.1.2.md"
PUBLIC_PROTOCOL = PUBLIC_COHERENCE_DIR / "Polymyth_Coherence_AI_Application_Protocol_V5.1.2.md"
SCHEMA = COHERENCE_DIR / "Polymyth_Coherence_Assessment_Schema_V5.1.2.json"
PUBLIC_SCHEMA = PUBLIC_COHERENCE_DIR / "Polymyth_Coherence_Assessment_Schema_V5.1.2.json"
COHERENCE_PAGE = COHERENCE_DIR / "index.html"
PUBLIC_COHERENCE_PAGE = PUBLIC_COHERENCE_DIR / "index.html"
COHERENCE_ADDENDUM = ROOT / "polymyth" / "methodologylist" / "polymyth-coherence-routing-addendum.js"
PUBLIC_COHERENCE_ADDENDUM = ROOT / "public" / "polymyth" / "methodologylist" / "polymyth-coherence-routing-addendum.js"
EDITABLE_INSTRUMENT = EDITABLE_DIR / "Polymyth_Coherence_Assessment_Instrument.xlsx"
WORKED_APPLICATIONS = EDITABLE_DIR / "Polymyth_Coherence_Worked_Applications.xlsx"
WORKED_APPLICATIONS_NAME = "Polymyth_Coherence_Worked_Applications.xlsx"

EXPECTED_FILE_SHA256 = "0624c76e0ae1351dcfe6a9d2cabf6e4e581820b63a0fc40e5ad755bbd1d93550"
EXPECTED_FORMULA_COUNT = 85768
EXPECTED_FORMULA_SIGNATURE = "94a9cd749e53b71c056c2a41c68d27cdb9a344eaf9bb3865b37a27f2b58684be"

EXPECTED_SHEETS = [
    "START HERE",
    "GOVERNING RULES",
    "INTERNAL",
    "MEZO",
    "EXTERNAL",
    "CROSS-LAYER",
    "BOUNDARY CASES",
    "DEBATES",
    "SOURCES",
    "RULE-SOURCE MAP",
    "COVERAGE",
    "CODEBOOK",
    "CHANGELOG",
    "AUDIT INDEX",
    "AUDIT PROTOCOL",
    "SCENE-EVENT LEDGER",
    "STATE-TIME-SPACE",
    "RULE-CAPABILITY",
    "KNOWLEDGE-DECISION",
    "CAUSAL-TRANSITIONS",
    "FINDINGS-REPAIRS",
    "RED-TEAM QA",
    "REGRESSION TESTS",
    "QA MASTER",
    "MEZO RELATIONS",
    "EXTERNAL CLAIMS",
    "SCOPE-CORPUS",
    "PRODUCTION DEFECTS",
    "ORGANIZING FORM",
    "TRUTH-DIRECTION",
    "SECOND-ORDER",
    "INTERPRETATIONS",
    "V5 INTEGRITY",
    "QA EVIDENCE",
    "SALIENCE PROFILE",
]

APPLICATION_SHEETS = [
    "AUDIT INDEX",
    "SCENE-EVENT LEDGER",
    "STATE-TIME-SPACE",
    "RULE-CAPABILITY",
    "KNOWLEDGE-DECISION",
    "CAUSAL-TRANSITIONS",
    "FINDINGS-REPAIRS",
    "RED-TEAM QA",
    "MEZO RELATIONS",
    "EXTERNAL CLAIMS",
    "SCOPE-CORPUS",
    "PRODUCTION DEFECTS",
    "ORGANIZING FORM",
    "TRUTH-DIRECTION",
    "SECOND-ORDER",
    "INTERPRETATIONS",
    "QA EVIDENCE",
    "SALIENCE PROFILE",
]


def fail(message: str) -> None:
    raise SystemExit(f"POLYMYTH COHERENCE WORKBOOK FAILED — {message}")


def require(condition: bool, message: str) -> None:
    if not condition:
        fail(message)


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def first_column_ids(workbook, sheet_name: str, pattern: str) -> list[str]:
    matcher = re.compile(pattern)
    values: list[str] = []
    for row in workbook[sheet_name].iter_rows(min_row=7, max_col=1):
        value = row[0].value
        if isinstance(value, str) and matcher.fullmatch(value):
            values.append(value)
    return values


def row_by_id(workbook, sheet_name: str, row_id: str):
    for row in workbook[sheet_name].iter_rows(min_row=7):
        if row[0].value == row_id:
            return row
    fail(f"{sheet_name} is missing {row_id}")


def require_text_tokens(path: Path, tokens: tuple[str, ...]) -> str:
    text = path.read_text(encoding="utf-8")
    for token in tokens:
        require(token in text, f"{path.relative_to(ROOT)} is missing {token}")
    return text


def is_blank_case_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, list):
        return not value
    if isinstance(value, dict):
        return all(is_blank_case_value(item) for item in value.values())
    return False


if SITE_ONLY:
    # A hosted build has only SITE_PACKAGE. Every byte read in this branch is
    # deliberately rooted at ROOT; private release masters are not consulted.
    deploy_pairs = (
        (COHERENCE_PAGE, PUBLIC_COHERENCE_PAGE),
        (INSTRUMENT, PUBLIC_INSTRUMENT),
        (PROTOCOL, PUBLIC_PROTOCOL),
        (SCHEMA, PUBLIC_SCHEMA),
        (COHERENCE_ADDENDUM, PUBLIC_COHERENCE_ADDENDUM),
    )
    for source_path, public_path in deploy_pairs:
        require(source_path.is_file(), f"missing {source_path.relative_to(ROOT)}")
        require(public_path.is_file(), f"missing {public_path.relative_to(ROOT)}")
        require(
            source_path.read_bytes() == public_path.read_bytes(),
            f"source/public Coherence asset differs: {source_path.relative_to(ROOT)}",
        )

    coherence_html = require_text_tokens(
        COHERENCE_PAGE,
        (
            "https://seminarschools.com/polymyth/coherence/",
            'data-route-type="archive"',
            'data-shared-geometry-exempt="star-file"',
            'data-star-file-page="true"',
            "Polymyth_Coherence_Assessment_Instrument_V5.1.2.xlsx",
            "Polymyth_Coherence_AI_Application_Protocol_V5.1.2.md",
            "Polymyth_Coherence_Assessment_Schema_V5.1.2.json",
            "Download the blank instrument",
            "There are no completed case audits in these downloads.",
            "Applicable with Adaptation",
            "Outside Protocol Scope",
            "Screening can locate candidates and defects",
            "AI assists; people decide.",
            "genuinely different independent reviewer",
            "Internal is the absolute gate.",
            "Looking for the wider method?",
        ),
    )
    require(
        "V5.1.1" not in coherence_html
        and "Polymyth_Coherence_V5.1.1.xlsx" not in coherence_html,
        "source Polymyth Coherence page still exposes V5.1.1",
    )
    require_text_tokens(
        PROTOCOL,
        (
            "no completed case audit",
            "Applicable with Adaptation",
            "Outside Protocol Scope",
            "Screening may locate candidates",
            "Never use `Polymyth_Coherence_Worked_Applications.xlsx` or a prior case result as an answer key.",
            "A human adjudicator must verify the evidence",
            "genuinely different independent reviewer",
            "Automated output alone cannot establish",
        ),
    )
    try:
        schema = json.loads(SCHEMA.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"{SCHEMA.relative_to(ROOT)} is invalid JSON: {exc}")
    require(schema.get("$schema") == "https://json-schema.org/draft/2020-12/schema", "blank JSON schema does not declare Draft 2020-12")
    require(
        schema.get("$ref") == "#/$defs/caseHandoff"
        and isinstance(schema.get("$defs", {}).get("caseHandoff"), dict),
        "blank JSON schema does not expose its case handoff through top-level $ref/$defs",
    )
    instrument_meta = schema.get("x-instrument", {})
    require(instrument_meta.get("version") == "V5.1.2", "blank JSON schema has the wrong instrument version")
    require(instrument_meta.get("sha256") == EXPECTED_FILE_SHA256, "blank JSON schema has the wrong workbook hash")
    blank_state = schema.get("x-blankState", {})
    require(blank_state.get("completedCaseCount") == 0, "blank JSON schema claims a completed case")
    require(blank_state.get("templateIsACompletedCase") is False, "blank JSON schema treats TEMPLATE-001 as completed")
    examples = schema.get("examples")
    blank_case = examples[0] if isinstance(examples, list) and examples else None
    require(
        isinstance(blank_case, dict)
        and all(is_blank_case_value(value) for value in blank_case.values()),
        "blank JSON schema contains a case answer",
    )


workbook_paths = (INSTRUMENT,) if SITE_ONLY else (INSTRUMENT, EDITABLE_INSTRUMENT)
for workbook_path in workbook_paths:
    require(workbook_path.is_file(), f"missing {workbook_path.relative_to(PACKAGE_ROOT)}")
    actual_hash = file_sha256(workbook_path)
    require(
        actual_hash == EXPECTED_FILE_SHA256,
        f"{workbook_path.relative_to(PACKAGE_ROOT)} hash is {actual_hash}, expected {EXPECTED_FILE_SHA256}",
    )
if not SITE_ONLY:
    require(WORKED_APPLICATIONS.is_file(), f"missing {WORKED_APPLICATIONS.relative_to(PACKAGE_ROOT)}")
require(
    not any(ROOT.rglob(WORKED_APPLICATIONS_NAME)),
    "worked applications workbook must remain outside SITE_PACKAGE and its public deploy tree",
)

obsolete_paths = [
    COHERENCE_DIR / "Polymyth_Coherence_V5.1.1.xlsx",
    ROOT / "public" / "polymyth" / "coherence" / "Polymyth_Coherence_V5.1.1.xlsx",
]
if not SITE_ONLY:
    obsolete_paths.append(EDITABLE_DIR / "Polymyth_Coherence.xlsx")
for obsolete in obsolete_paths:
    require(not obsolete.exists(), f"obsolete workbook remains: {obsolete.relative_to(PACKAGE_ROOT)}")

workbook = load_workbook(INSTRUMENT, read_only=False, data_only=False)
require(workbook.sheetnames == EXPECTED_SHEETS, "the exact 35-sheet order changed")

formula_lines: list[str] = []
reference_errors: list[str] = []
for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formula_lines.append(f"{sheet.title}\t{cell.coordinate}\t{value}\n")
                if "#REF!" in value:
                    reference_errors.append(f"{sheet.title}!{cell.coordinate}")

formula_bytes = "".join(formula_lines).encode("utf-8")
formula_signature = hashlib.sha256(formula_bytes).hexdigest()
require(len(formula_lines) == EXPECTED_FORMULA_COUNT, f"formula count is {len(formula_lines)}")
require(
    formula_signature == EXPECTED_FORMULA_SIGNATURE,
    f"formula signature is {formula_signature}, expected {EXPECTED_FORMULA_SIGNATURE}",
)
require(not reference_errors, f"formula #REF! errors found at {', '.join(reference_errors[:8])}")


def formula_records(candidate_workbook) -> list[str]:
    records: list[str] = []
    for sheet in candidate_workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    records.append(f"{sheet.title}\t{cell.coordinate}\t{value}\n")
    return records


def populated_cell_records(worksheet, max_row: int | None = None) -> list[tuple[str, str, object]]:
    records: list[tuple[str, str, object]] = []
    row_limit = worksheet.max_row if max_row is None else max_row
    for row in worksheet.iter_rows(min_row=1, max_row=row_limit):
        for cell in row:
            if cell.value not in (None, ""):
                records.append((cell.coordinate, cell.data_type, cell.value))
    return records


def require_complete_validation_ranges(candidate, label: str) -> None:
    external_validation = {
        str(item.sqref) for item in candidate["EXTERNAL"].data_validations.dataValidation
    }
    external_conditional = {
        str(item.sqref) for item in candidate["EXTERNAL"].conditional_formatting
    }
    source_validation = {
        str(item.sqref) for item in candidate["SOURCES"].data_validations.dataValidation
    }
    require("N7:N42" in external_validation, f"{label} omits E-036 from EXTERNAL validation")
    require("N7:N42" in external_conditional, f"{label} omits E-036 from EXTERNAL formatting")
    require("P7:P292" in source_validation, f"{label} omits S-276 through S-286 from SOURCES validation")
    require("N7:N41" not in external_validation | external_conditional, f"{label} retains truncated EXTERNAL ranges")
    require("P7:P281" not in source_validation, f"{label} retains the truncated SOURCES range")


require_complete_validation_ranges(workbook, "canonical instrument")
if not SITE_ONLY:
    worked = load_workbook(WORKED_APPLICATIONS, read_only=False, data_only=False)
    require(
        worked.sheetnames == EXPECTED_SHEETS,
        "worked applications sheet order changed",
    )
    require_complete_validation_ranges(worked, "worked applications")

    # A case file may change only case-input values. The complete formula map
    # stays byte-for-byte equivalent at the formula-record level, including
    # application sheets, so a case cannot overwrite or invent calculation logic.
    worked_formula_lines = formula_records(worked)
    require(
        worked_formula_lines == formula_lines,
        "worked applications formula coordinates or formulas changed",
    )

    # Criteria, governance, protocol, reference, and instrument-test sheets
    # remain value-identical to the canonical instrument. Application sheets
    # retain their six-row structural headers and preallocated bounds while rows
    # 7 onward may contain case evidence, rulings, review, and results.
    immutable_sheets = [name for name in EXPECTED_SHEETS if name not in APPLICATION_SHEETS]
    for sheet_name in immutable_sheets:
        require(
            populated_cell_records(worked[sheet_name]) == populated_cell_records(workbook[sheet_name]),
            f"worked applications changed immutable sheet {sheet_name}",
        )
    for sheet_name in APPLICATION_SHEETS:
        canonical_sheet = workbook[sheet_name]
        worked_sheet = worked[sheet_name]
        require(
            populated_cell_records(worked_sheet, 6) == populated_cell_records(canonical_sheet, 6),
            f"worked applications changed the structural header of {sheet_name}",
        )
        require(
            worked_sheet.max_row == canonical_sheet.max_row
            and worked_sheet.max_column == canonical_sheet.max_column,
            f"worked applications wrote outside the preallocated bounds of {sheet_name}",
        )

expected_id_sets = {
    "GOVERNING RULES": [f"G-{number:03d}" for number in range(1, 53)],
    "INTERNAL": [f"I-{number:03d}" for number in range(1, 178)],
    "MEZO": [f"M-{number:03d}" for number in range(1, 60)],
    "EXTERNAL": [f"E-{number:03d}" for number in range(1, 37)],
    "CROSS-LAYER": [f"X-{number:03d}" for number in range(1, 46)],
    "DEBATES": [f"D-{number:03d}" for number in range(1, 90)],
    "SOURCES": [f"S-{number:03d}" for number in range(1, 287)],
    "RULE-SOURCE MAP": [f"RS-{number:04d}" for number in range(1, 845)],
    "QA MASTER": [f"Q-{number:03d}" for number in range(1, 41)],
}
id_patterns = {
    "GOVERNING RULES": r"G-\d{3}",
    "INTERNAL": r"I-\d{3}",
    "MEZO": r"M-\d{3}",
    "EXTERNAL": r"E-\d{3}",
    "CROSS-LAYER": r"X-\d{3}",
    "DEBATES": r"D-\d{3}",
    "SOURCES": r"S-\d{3}",
    "RULE-SOURCE MAP": r"RS-\d{4}",
    "QA MASTER": r"Q-\d{3}",
}
for sheet_name, expected_ids in expected_id_sets.items():
    actual_ids = first_column_ids(workbook, sheet_name, id_patterns[sheet_name])
    require(actual_ids == expected_ids, f"{sheet_name} IDs or count changed")

protocol_ids = first_column_ids(workbook, "AUDIT PROTOCOL", r"P(?:[0-9]|1[0-9]|2[0-3])")
require(protocol_ids == [f"P{number}" for number in range(24)], "AUDIT PROTOCOL must contain P0 through P23")

require(
    workbook["START HERE"]["A1"].value == "Polymyth Coherence: Internal, Mezo, and External Logic, V5.1.2",
    "START HERE version changed",
)
require("V5.1.2 critique integration" in str(workbook["START HERE"]["B24"].value), "START HERE release note changed")

g052 = row_by_id(workbook, "GOVERNING RULES", "G-052")
require(g052[4].value == "Ratified", "G-052 is not Ratified")
require("Automated output alone cannot establish" in str(g052[2].value), "G-052 automation boundary changed")
row_by_id(workbook, "GOVERNING RULES", "G-033")
row_by_id(workbook, "GOVERNING RULES", "G-044")
row_by_id(workbook, "INTERNAL", "I-146")

m054 = row_by_id(workbook, "MEZO", "M-054")
m059 = row_by_id(workbook, "MEZO", "M-059")
require(m054[13].value == "Superseded", "M-054 must remain Superseded")
require(m059[13].value == "Active", "M-059 must remain Active")
require("least-distorting, explicitly situated bridge" in str(m059[3].value), "M-059 situated bridge wording changed")

for debate_id in ("D-082", "D-083"):
    require(row_by_id(workbook, "DEBATES", debate_id)[12].value == "Open", f"{debate_id} must remain Open")
for debate_id in ("D-084", "D-085", "D-086", "D-087", "D-088", "D-089"):
    require(row_by_id(workbook, "DEBATES", debate_id)[12].value == "Resolved", f"{debate_id} must remain Resolved")

p0 = row_by_id(workbook, "AUDIT PROTOCOL", "P0")
p23 = row_by_id(workbook, "AUDIT PROTOCOL", "P23")
require("Applicable with Adaptation" in str(p0[2].value), "P0 protocol-fit states changed")
require("Outside Protocol Scope" in str(p0[2].value), "P0 outside-scope boundary changed")
require("Screening may locate candidates" in str(p23[2].value), "P23 screening boundary changed")
require("V5.1.2" == row_by_id(workbook, "CHANGELOG", "V5.1.2")[0].value, "V5.1.2 changelog row missing")

audit_ids = [
    workbook["AUDIT INDEX"].cell(row=row, column=1).value
    for row in range(7, workbook["AUDIT INDEX"].max_row + 1)
    if workbook["AUDIT INDEX"].cell(row=row, column=1).value not in (None, "")
]
require(audit_ids == ["TEMPLATE-001"], "instrument contains a non-template Audit Index case")
require("Template only" in str(workbook["AUDIT INDEX"]["O7"].value), "template warning changed")

for sheet_name in APPLICATION_SHEETS:
    if sheet_name in ("AUDIT INDEX", "RED-TEAM QA"):
        continue
    case_ids = [
        workbook[sheet_name].cell(row=row, column=1).value
        for row in range(7, workbook[sheet_name].max_row + 1)
        if workbook[sheet_name].cell(row=row, column=1).value not in (None, "")
    ]
    require(not case_ids, f"{sheet_name} contains case data: {case_ids[:3]}")

red_team = workbook["RED-TEAM QA"]
require(
    [red_team.cell(row=row, column=1).value for row in range(7, 47)] == ["TEMPLATE-001"] * 40,
    "RED-TEAM QA template Case IDs changed",
)
require(
    [red_team.cell(row=row, column=2).value for row in range(7, 47)]
    == [f"Q-{number:03d}" for number in range(1, 41)],
    "RED-TEAM QA must contain Q-001 through Q-040 in order",
)
require(
    [red_team.cell(row=row, column=6).value for row in range(7, 47)] == ["Not checked"] * 40,
    "RED-TEAM QA template must remain unchecked",
)
require(
    not any(red_team.cell(row=row, column=1).value not in (None, "") for row in range(47, red_team.max_row + 1)),
    "RED-TEAM QA contains non-template case rows",
)

cached = load_workbook(INSTRUMENT, read_only=False, data_only=True)
require(cached["V5 INTEGRITY"]["H4"].value == "PASS", "cached V5 INTEGRITY result is not PASS")
require(
    [cached["QA MASTER"].cell(row=row, column=8).value for row in range(7, 47)] == ["PASS"] * 40,
    "one or more cached QA MASTER checks fail",
)
require(
    [cached["REGRESSION TESTS"].cell(row=row, column=18).value for row in range(7, 25)] == ["PASS"] * 18,
    "one or more executable gate regressions fail",
)
require(
    [cached["REGRESSION TESTS"].cell(row=row, column=8).value for row in range(28, 32)] == ["PASS"] * 4,
    "one or more truth-direction regressions fail",
)

cached_errors: list[str] = []
for sheet in cached.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "e":
                cached_errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
require(not cached_errors, f"cached spreadsheet errors found: {', '.join(cached_errors[:8])}")

if SITE_ONLY:
    print(
        "POLYMYTH COHERENCE WORKBOOK SITE-ONLY PASSED — "
        "source/public deploy parity, byte-exact blank V5.1.2 workbook, protocol, answer-free schema, "
        "routing addendum, 35 sheets, 85,768 formulas, 24 protocol entries, 40 QA checks, "
        "22 regressions, zero completed cases, and no worked-applications leak verified "
        "without private release masters."
    )
else:
    print(
        "POLYMYTH COHERENCE WORKBOOK PASSED — "
        "two byte-exact blank V5.1.2 instruments plus one structurally locked, editable applications workbook; "
        "35 sheets, 85,768 formulas, 24 protocol entries, 40 QA checks, 22 regressions, "
        "and zero completed cases in the public instrument verified."
    )
