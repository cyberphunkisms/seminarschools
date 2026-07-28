#!/usr/bin/env python3
"""Preserve the superseded 90/38/1 #MeToo premise audit row by row."""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook


SPLIT_OUTSIDE = "EXCOMMUNICABLE / OUTSIDE"
SPLIT_INTERNAL = "INTERNAL GORGONWARS"
SPLIT_UNRESOLVED = "UNRESOLVED"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def derive_split(pip: str, standpoint: str) -> str:
    if pip == "Rejects" or standpoint == "Rejects":
        return SPLIT_OUTSIDE
    if pip == "Retains" and standpoint == "Retains":
        return SPLIT_INTERNAL
    return SPLIT_UNRESOLVED


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(
            "usage: extract-withdrawn-metoo-premise-audit.py INPUT.xlsx OUTPUT_DIR"
        )

    source = Path(sys.argv[1]).resolve()
    output_dir = Path(sys.argv[2]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = "metoo-critical-academic-premise-audit-WITHDRAWN-90-38-1-2026-07-27"
    workbook_path = output_dir / f"{stem}.xlsx"
    jsonl_path = output_dir / f"{stem}.jsonl"

    workbook = load_workbook(source, read_only=False, data_only=False)
    sheet = workbook["System split"]
    records = []
    for row in sheet.iter_rows(min_row=13, max_row=141, values_only=True):
        record_id = str(row[0] or "").strip()
        if not record_id:
            continue
        pip = str(row[2] or "").strip()
        standpoint = str(row[3] or "").strip()
        records.append(
            {
                "withdrawn_audit": {
                    "version": "initial-premise-split-2026-07-27",
                    "status": "WITHDRAWN",
                    "reason": (
                        "The converter treated use of testimony, personal narrative, "
                        "situated knowledge, or related evidence as premise retention "
                        "without requiring article-level premise evidence."
                    ),
                },
                "id": record_id,
                "split": derive_split(pip, standpoint),
                "pip": pip,
                "standpoint": standpoint,
                "confidence": str(row[4] or "").strip(),
                "evidence": str(row[5] or "").strip(),
                "evidenceUrls": str(row[6] or "").strip(),
                "originalStance": str(row[7] or "").strip(),
                "theme": str(row[8] or "").strip(),
                "year": row[9],
                "authors": str(row[10] or "").strip(),
                "title": str(row[11] or "").strip(),
                "journal": str(row[12] or "").strip(),
                "doi": str(row[13] or "").strip(),
                "stableUrl": str(row[14] or "").strip(),
            }
        )

    counts: dict[str, int] = {}
    for record in records:
        counts[record["split"]] = counts.get(record["split"], 0) + 1
    expected = {
        SPLIT_OUTSIDE: 1,
        SPLIT_INTERNAL: 90,
        SPLIT_UNRESOLVED: 38,
    }
    if len(records) != 129 or counts != expected:
        raise RuntimeError(
            f"withdrawn audit mismatch: rows={len(records)} counts={counts}"
        )

    shutil.copyfile(source, workbook_path)
    jsonl_path.write_text(
        "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records),
        encoding="utf-8",
    )
    for artifact in (workbook_path, jsonl_path):
        artifact.with_suffix(artifact.suffix + ".sha256").write_text(
            f"{digest(artifact)}  {artifact.name}\n",
            encoding="utf-8",
        )

    print(
        json.dumps(
            {
                "rows": len(records),
                "counts": counts,
                "workbook": str(workbook_path),
                "jsonl": str(jsonl_path),
                "workbook_sha256": digest(workbook_path),
                "jsonl_sha256": digest(jsonl_path),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
