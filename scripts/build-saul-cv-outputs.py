#!/usr/bin/env python3
"""Build every Saul CV document from one editable JSON source.

Edit data/saul-ultimate-school-cv-2026.json, then run:
    npm run build:saul-cv

The builder produces:
* Gmail and ProtonMail one-page application DOCX/PDF files
* one-page role-focused designed/ATS PDFs and text exports
* the complete EVERYTHING PDF
* a ZIP containing every CV output
* a machine-readable output manifest
"""

from __future__ import annotations

import hashlib
import html
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    KeepTogether,
    Paragraph as RLParagraph,
    SimpleDocTemplate,
    Spacer,
)


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "saul-ultimate-school-cv-2026.json"
TEMPLATE_PATH = ROOT / "templates" / "saul-ultimate-school-cv-template.docx"
DOWNLOADS = ROOT / "saul" / "downloads"
MANIFEST_PATH = ROOT / "data" / "saul-cv-pdf-manifest.json"
HISTORICAL_PATH = ROOT / "data" / "saul-cv-records.json"

INK = HexColor("#182126")
MUTED = HexColor("#596166")
RULE = HexColor("#B1B5B6")
PALE = HexColor("#F0F1F1")
PAGE_W, PAGE_H = letter

FONT_ROOT = Path(
    "/opt/codex/runtimes/codex-primary-runtime/dependencies/native/"
    "libreoffice-headless/libreoffice/share/fonts/truetype"
)


def load_data() -> dict:
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))


def all_application_records(data: dict) -> list[dict]:
    return [
        record
        for section in data["experience_sections"]
        for record in section["records"]
    ]


def section_counts(data: dict) -> list[int]:
    return [len(section["records"]) for section in data["experience_sections"]]


def set_run_font(
    run,
    name: str,
    size: float,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_run_scale(run, scale: int) -> None:
    if scale >= 100:
        return
    rpr = run._element.get_or_add_rPr()
    width = rpr.find(qn("w:w"))
    if width is None:
        width = OxmlElement("w:w")
        rpr.append(width)
    width.set(qn("w:val"), str(scale))


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_hyperlink(
    paragraph: Paragraph,
    text: str,
    url: str,
    *,
    font: str,
    size: float,
    bold: bool = False,
) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font)
    rpr.append(fonts)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(size_el)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(size_cs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_plain_paragraph(
    paragraph: Paragraph,
    text: str,
    *,
    font: str,
    size: float,
    bold: bool = False,
    italic: bool = False,
) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold=bold, italic=italic)


def split_organization(value: str) -> tuple[str, str]:
    if value.endswith(")") and " (" in value:
        head, tail = value.rsplit(" (", 1)
        return head, f" ({tail}"
    return value, ""


def clone_paragraph_after(source: Paragraph, after: Paragraph) -> Paragraph:
    clone = deepcopy(source._p)
    for child in list(clone):
        if child.tag != qn("w:pPr"):
            clone.remove(child)
    after._p.addnext(clone)
    return Paragraph(clone, after._parent)


def build_experience_row(
    paragraph: Paragraph,
    record: dict,
    *,
    row_leading: float,
    character_scale: int,
) -> None:
    clear_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(row_leading)

    role = paragraph.add_run(record["role"])
    set_run_font(role, "Calibri", 11, bold=True, italic=False)
    set_run_scale(role, character_scale)
    separator = paragraph.add_run(" | ")
    set_run_font(separator, "Calibri", 11, bold=False, italic=False)
    set_run_scale(separator, character_scale)
    description = paragraph.add_run(record["description"])
    set_run_font(description, "Calibri", 10, bold=False, italic=False)
    set_run_scale(description, character_scale)
    separator = paragraph.add_run(" | ")
    set_run_font(separator, "Calibri", 11, bold=False, italic=False)
    set_run_scale(separator, character_scale)
    organization, location = split_organization(record["organization"])
    org_run = paragraph.add_run(organization)
    set_run_font(org_run, "Calibri", 11, bold=False, italic=True)
    set_run_scale(org_run, character_scale)
    if location:
        location_run = paragraph.add_run(location)
        set_run_font(location_run, "Calibri", 11, bold=False, italic=False)
        set_run_scale(location_run, character_scale)
    tab = paragraph.add_run("\t")
    set_run_font(tab, "Calibri", 11, bold=False, italic=False)
    dates = paragraph.add_run(record["dates"])
    set_run_font(dates, "Calibri", 11, bold=False, italic=False)


def replace_experience_sections(doc: Document, data: dict, row_leading: float) -> None:
    closing_title = "EDUCATION, PROFESSIONAL LEARNING & LANGUAGES"
    boundaries = [
        section["title"].upper() for section in data["experience_sections"]
    ] + [closing_title]
    sample = next(
        p for p in doc.paragraphs if p.text.startswith("Occasional Instructor |")
    )
    scale_by_body: dict[str, int] = {}
    for paragraph in doc.paragraphs:
        if "\t" not in paragraph.text:
            continue
        body = paragraph.text.split("\t", 1)[0]
        scale = 100
        for run in paragraph.runs:
            rpr = run._element.rPr
            width = rpr.find(qn("w:w")) if rpr is not None else None
            if width is not None:
                scale = int(width.get(qn("w:val"), "100"))
                break
        scale_by_body[body] = scale

    for index in range(len(data["experience_sections"]) - 1, -1, -1):
        section = data["experience_sections"][index]
        paragraphs = doc.paragraphs
        heading_index = next(
            i for i, p in enumerate(paragraphs) if p.text == section["title"].upper()
        )
        next_index = next(
            i
            for i, p in enumerate(paragraphs)
            if i > heading_index and p.text == boundaries[index + 1]
        )
        spacer = paragraphs[heading_index + 1]
        for paragraph in paragraphs[heading_index + 2 : next_index]:
            paragraph._element.getparent().remove(paragraph._element)

        cursor = spacer
        for record in section["records"]:
            cursor = clone_paragraph_after(sample, cursor)
            body = (
                f"{record['role']} | {record['description']} | "
                f"{record['organization']}"
            )
            build_experience_row(
                cursor,
                record,
                row_leading=row_leading,
                character_scale=scale_by_body.get(body, 100),
            )


def set_contact_line(paragraph: Paragraph, data: dict, email: str) -> None:
    contact = data["contact"]
    clear_paragraph(paragraph)
    run = paragraph.add_run(contact["location"])
    set_run_font(run, "Comfortaa", 9)
    run = paragraph.add_run("\t| ")
    set_run_font(run, "Comfortaa", 9)
    add_hyperlink(
        paragraph,
        contact["phone"],
        "tel:+14167710382",
        font="Comfortaa",
        size=9,
    )
    run = paragraph.add_run("\t| ")
    set_run_font(run, "Comfortaa", 9)
    add_hyperlink(
        paragraph,
        email,
        f"mailto:{email}",
        font="Comfortaa",
        size=9,
    )
    run = paragraph.add_run("\t| ")
    set_run_font(run, "Comfortaa", 9)
    add_hyperlink(
        paragraph,
        contact["site_label"],
        contact["site_url"],
        font="Comfortaa",
        size=9,
    )


def set_footer_line(paragraph: Paragraph, data: dict) -> None:
    contact = data["contact"]
    clear_paragraph(paragraph)
    run = paragraph.add_run("Reviews  ")
    set_run_font(run, "Lexend Exa", 8, bold=True)
    add_hyperlink(
        paragraph,
        contact["reviews_label"],
        contact["reviews_url"],
        font="Lexend Exa",
        size=8,
        bold=True,
    )
    run = paragraph.add_run("  |  References available on request")
    set_run_font(run, "Lexend Exa", 8, bold=True)


def set_education_line(paragraph: Paragraph, entries: list[str]) -> None:
    clear_paragraph(paragraph)
    for index, entry in enumerate(entries):
        if index:
            separator = paragraph.add_run("  |  ")
            set_run_font(separator, "Calibri", 8.5)
        if ", " in entry:
            degree, remainder = entry.split(", ", 1)
            degree_run = paragraph.add_run(degree)
            set_run_font(degree_run, "Calibri", 9, bold=True, italic=True)
            rest_run = paragraph.add_run(", " + remainder)
            set_run_font(rest_run, "Calibri", 8.5, italic=True)
        else:
            run = paragraph.add_run(entry)
            set_run_font(run, "Calibri", 8.5, italic=True)


def set_labeled_line(paragraph: Paragraph, label: str, values: list[str]) -> None:
    clear_paragraph(paragraph)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, "Cambria", 8, bold=True)
    tab = paragraph.add_run("\t")
    set_run_font(tab, "Cambria", 7)
    values_run = paragraph.add_run(" | ".join(values))
    set_run_font(values_run, "Cambria", 7)


def application_courses(data: dict) -> list[str]:
    codes = ", ".join(code for code, _ in data["courses"]["ontario"])
    international_map = {
        "IB Diploma Courses": "IB Diploma",
        "IB Theory of Knowledge": "IB TOK",
        "AP English Language & Composition": "AP English",
        "IELTS Preparation": "IELTS",
        "ESL across CEFR A1–C1": "ESL A1–C1",
    }
    international = []
    for item in data["courses"]["international"]:
        if item in {
            "Academic English Transition Support",
            "University Preparation & Admissions Support",
            "Early Childhood STEM, Drama & English",
        }:
            continue
        international.append(international_map.get(item, item))
    return [codes, " | ".join(international)]


def build_application_docx(
    data: dict,
    output: Path,
    email: str,
    *,
    row_leading: float,
) -> None:
    if not TEMPLATE_PATH.exists():
        raise RuntimeError(f"Missing editable CV template: {TEMPLATE_PATH}")
    doc = Document(TEMPLATE_PATH)
    paragraphs = doc.paragraphs

    set_plain_paragraph(
        paragraphs[0],
        data["contact"]["name"],
        font="Comfortaa",
        size=21,
        bold=True,
    )
    set_plain_paragraph(
        paragraphs[1],
        data["contact"]["role"],
        font="Calibri",
        size=13,
    )
    set_contact_line(paragraphs[2], data, email)
    set_plain_paragraph(
        paragraphs[3],
        data["profile"],
        font="Calibri",
        size=9.5,
    )
    clear_paragraph(paragraphs[4])
    skill_groups = [
        data["core_skills"][0:4],
        data["core_skills"][4:8],
        data["core_skills"][8:12],
    ]
    for index, group in enumerate(skill_groups):
        if index:
            paragraphs[4].add_run("\n")
        prefix = "CORE SKILLS | " if index == 0 else ""
        run = paragraphs[4].add_run(prefix + " | ".join(group))
        set_run_font(run, "Cambria", 7.5, bold=index == 0)

    replace_experience_sections(doc, data, row_leading)

    paragraphs = doc.paragraphs
    closing_index = next(
        i
        for i, paragraph in enumerate(paragraphs)
        if paragraph.text == "EDUCATION, PROFESSIONAL LEARNING & LANGUAGES"
    )
    set_education_line(paragraphs[closing_index + 2], data["education"][:2])
    set_education_line(paragraphs[closing_index + 3], data["education"][2:3])
    details = [
        ("Credentials", data["credentials"]),
        ("Professional Learning", data["professional_learning"]),
        ("Methods & Tools", data["methods_tools"]),
        ("Courses", application_courses(data)),
        ("Languages", data["languages"]),
    ]
    for offset, (label, values) in enumerate(details, start=4):
        set_labeled_line(paragraphs[closing_index + offset], label, values)
    set_footer_line(paragraphs[closing_index + 9], data)

    doc.core_properties.title = "Saul Karim Nassau | School and Teaching CV"
    doc.core_properties.subject = (
        "One-page educator, research, community & employment résumé"
    )
    doc.core_properties.author = "Saul Karim Nassau"
    doc.core_properties.keywords = (
        "teaching, education, OSSD, IB, AP, ESL, research, community, Greenpeace"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def clone_docx_email_only(source: Path, target: Path, old: str, new: str) -> None:
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(
        target, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename.endswith(".xml") or info.filename.endswith(".rels"):
                payload = payload.replace(old.encode("utf-8"), new.encode("utf-8"))
            clone = zipfile.ZipInfo(info.filename, (2026, 7, 28, 0, 0, 0))
            clone.compress_type = info.compress_type
            clone.external_attr = info.external_attr
            clone.create_system = info.create_system
            dst.writestr(clone, payload)


def convert_docx_to_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    if pdf_path.exists():
        pdf_path.unlink()
    with tempfile.TemporaryDirectory(prefix="saul-cv-lo-") as temp:
        temp_path = Path(temp)
        profile = temp_path / "profile"
        home = temp_path / "home"
        profile.mkdir()
        home.mkdir()
        env = os.environ.copy()
        env["HOME"] = str(home)
        command = [
            shutil.which("soffice") or "soffice",
            "--headless",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ]
        result = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            env=env,
        )
        if result.returncode or not pdf_path.exists():
            raise RuntimeError(
                f"LibreOffice PDF conversion failed for {docx_path}: "
                f"{result.stdout}\n{result.stderr}"
            )
    return pdf_path


def register_pdf_fonts() -> None:
    fonts = {
        "CVSans": FONT_ROOT / "NotoSans-Regular.ttf",
        "CVSans-Bold": FONT_ROOT / "NotoSans-Bold.ttf",
        "CVSans-Italic": FONT_ROOT / "NotoSans-Italic.ttf",
    }
    for name, path in fonts.items():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))


def wrap_lines(text: str, font: str, size: float, width: float) -> list[str]:
    words = str(text).split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else current + " " + word
        if pdfmetrics.stringWidth(candidate, font, size) <= width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_wrapped(
    cv: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    width: float,
    *,
    font: str,
    size: float,
    leading: float,
    color=INK,
    max_lines: int | None = None,
) -> float:
    lines = wrap_lines(text, font, size, width)
    if max_lines:
        lines = lines[:max_lines]
    cv.setFont(font, size)
    cv.setFillColor(color)
    for line in lines:
        cv.drawString(x, y, line)
        y -= leading
    return y


def records_for_focus(data: dict, focus_id: str) -> list[dict]:
    records = all_application_records(data)
    if focus_id == "general":
        return records
    return [record for record in records if focus_id in record.get("focus", [])]


def draw_compressed_text(
    cv: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    width: float,
    *,
    font: str,
    size: float,
    color=INK,
) -> None:
    natural = pdfmetrics.stringWidth(text, font, size)
    scale = min(100.0, max(72.0, width / max(natural, 1) * 100.0))
    obj = cv.beginText(x, y)
    obj.setFont(font, size)
    obj.setFillColor(color)
    obj.setHorizScale(scale)
    obj.textOut(text)
    cv.drawText(obj)


def modular_pdf(data: dict, module: dict, output: Path, *, ats: bool) -> None:
    records = records_for_focus(data, module["id"])
    if not records and module.get("archive_only"):
        records = all_application_records(data)
    cv = canvas.Canvas(str(output), pagesize=letter, pageCompression=1)
    cv.setTitle(
        f"Saul Karim Nassau - {module['label']} CV"
        + (" - ATS" if ats else "")
    )
    cv.setAuthor("Saul Karim Nassau")
    left = 38
    right = PAGE_W - 38
    width = right - left

    cv.setFillColor(INK)
    cv.setFont("CVSans-Bold", 18)
    cv.drawString(left, PAGE_H - 43, data["contact"]["name"])
    cv.setFont("CVSans-Bold", 9.5)
    cv.drawString(left, PAGE_H - 59, module["label"].upper())
    cv.setFillColor(MUTED)
    cv.setFont("CVSans", 7)
    contact = data["contact"]
    cv.drawRightString(
        right,
        PAGE_H - 43,
        f"{contact['location']} | {contact['phone']}",
    )
    cv.drawRightString(
        right,
        PAGE_H - 55,
        f"{contact['public_email']} | {contact['site_label']}",
    )
    cv.setStrokeColor(INK)
    cv.setLineWidth(1)
    cv.line(left, PAGE_H - 67, right, PAGE_H - 67)

    y = PAGE_H - 82
    y = draw_wrapped(
        cv,
        module["summary"],
        left,
        y,
        width,
        font="CVSans",
        size=7.7,
        leading=9.2,
        max_lines=3,
    )
    y -= 2
    skills = " | ".join(module.get("skills", data["core_skills"]))
    y = draw_wrapped(
        cv,
        "KEY SKILLS  |  " + skills,
        left,
        y,
        width,
        font="CVSans",
        size=6.8,
        leading=8.2,
        max_lines=3,
    )
    y -= 5
    cv.setFont("CVSans-Bold", 8.4)
    cv.setFillColor(INK)
    cv.drawString(left, y, f"SELECTED EXPERIENCE · {len(records)}")
    y -= 10

    bottom = 54
    available = max(1, y - bottom)
    row_leading = min(16.0, available / max(1, len(records)))
    row_leading = max(9.4, row_leading)
    date_width = 92
    content_width = width - date_width - 8
    role_size = 7.7 if ats else 8.0
    for record in records:
        row = (
            f"{record['role']} | {record['description']} | "
            f"{record['organization']}"
        )
        draw_compressed_text(
            cv,
            row,
            left,
            y,
            content_width,
            font="CVSans",
            size=role_size,
        )
        cv.setFont("CVSans", 7.5)
        cv.setFillColor(MUTED)
        cv.drawRightString(right, y, record["dates"])
        y -= row_leading

    cv.setStrokeColor(RULE)
    cv.setLineWidth(0.7)
    cv.line(left, 40, right, 40)
    cv.setFont("CVSans-Bold", 6.6)
    cv.setFillColor(MUTED)
    cv.drawCentredString(
        PAGE_W / 2,
        27,
        "REVIEWS  SEMINARSCHOOLS.COM/REVIEWS  |  REFERENCES AVAILABLE ON REQUEST",
    )
    cv.showPage()
    cv.save()


def modular_text(data: dict, module: dict) -> str:
    contact = data["contact"]
    records = records_for_focus(data, module["id"])
    if not records and module.get("archive_only"):
        records = all_application_records(data)
    lines = [
        contact["name"],
        module["label"].upper(),
        (
            f"{contact['location']} | {contact['phone']} | "
            f"{contact['public_email']} | {contact['site_label']}"
        ),
        "",
        "PROFILE",
        module["summary"],
        "",
        "KEY SKILLS",
        " | ".join(module.get("skills", data["core_skills"])),
        "",
        f"SELECTED EXPERIENCE ({len(records)})",
    ]
    lines.extend(
        (
            f"{record['role']} | {record['description']} | "
            f"{record['organization']} | {record['dates']}"
        )
        for record in records
    )
    lines.extend(
        [
            "",
            "EDUCATION",
            *data["education"],
            "",
            "CREDENTIALS",
            *data["credentials"],
            "",
            "LANGUAGES",
            *data["languages"],
            "",
            f"Reviews: {contact['reviews_label']}",
        ]
    )
    return "\n".join(lines) + "\n"


def historical_records(data: dict) -> list[dict]:
    payload = json.loads(HISTORICAL_PATH.read_text(encoding="utf-8"))
    records = list(payload.get("records", []))
    if not any("Greenpeace" in record.get("title", "") for record in records):
        greenpeace = next(
            record
            for record in all_application_records(data)
            if record["organization"] == "Greenpeace"
        )
        records.append(
            {
                "index": len(records),
                "section_weight": 1,
                "date": greenpeace["dates"],
                "categories": ["community", "volunteer"],
                "title": f"{greenpeace['role']}, Greenpeace",
                "note": greenpeace["description"],
                "url": None,
                "description": (
                    "Fundraising and volunteer coordination for Greenpeace."
                ),
                "status": "completed or ongoing experience",
                "source": (
                    "User-provided employment record and screenshot, 2026-07-28"
                ),
            }
        )
    return records


def everything_pdf(data: dict, output: Path) -> None:
    styles = {
        "title": ParagraphStyle(
            "title",
            fontName="CVSans-Bold",
            fontSize=20,
            leading=22,
            textColor=INK,
            spaceAfter=3,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            fontName="CVSans-Bold",
            fontSize=10.5,
            leading=12.5,
            textColor=INK,
            spaceAfter=5,
        ),
        "h2": ParagraphStyle(
            "h2",
            fontName="CVSans-Bold",
            fontSize=10,
            leading=12,
            textColor=INK,
            spaceBefore=8,
            spaceAfter=4,
        ),
        "h3": ParagraphStyle(
            "h3",
            fontName="CVSans-Bold",
            fontSize=8,
            leading=9.4,
            textColor=INK,
            spaceAfter=1,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="CVSans",
            fontSize=7.1,
            leading=8.8,
            textColor=INK,
            spaceAfter=2,
        ),
        "meta": ParagraphStyle(
            "meta",
            fontName="CVSans",
            fontSize=6.6,
            leading=8,
            textColor=MUTED,
            spaceAfter=2,
        ),
        "compact": ParagraphStyle(
            "compact",
            fontName="CVSans",
            fontSize=6.6,
            leading=8,
            textColor=INK,
            leftIndent=8,
            firstLineIndent=-8,
            spaceAfter=1,
        ),
    }

    def page_frame(cv: canvas.Canvas, doc) -> None:
        cv.saveState()
        cv.setFillColor(white)
        cv.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
        cv.setFillColor(INK)
        cv.rect(36, PAGE_H - 25, PAGE_W - 72, 3, fill=1, stroke=0)
        cv.setStrokeColor(RULE)
        cv.line(36, 29, PAGE_W - 36, 29)
        cv.setFont("CVSans", 6.2)
        cv.setFillColor(MUTED)
        cv.drawString(36, 18, "SEMINARSCHOOLS.COM/SAUL")
        cv.drawRightString(PAGE_W - 36, 18, str(doc.page))
        cv.restoreState()

    document = SimpleDocTemplate(
        str(output),
        pagesize=letter,
        leftMargin=38,
        rightMargin=38,
        topMargin=42,
        bottomMargin=38,
        title="Saul Karim Nassau - Complete Career Archive",
        author="Saul Karim Nassau",
    )
    contact = data["contact"]
    story = [
        RLParagraph(html.escape(contact["name"]), styles["title"]),
        RLParagraph("COMPLETE CAREER ARCHIVE", styles["subtitle"]),
        RLParagraph(
            html.escape(
                f"{contact['location']} | {contact['phone']} | "
                f"{contact['public_email']} | {contact['site_label']} | "
                f"{contact['reviews_label']}"
            ),
            styles["meta"],
        ),
        RLParagraph(html.escape(data["profile"]), styles["body"]),
        RLParagraph("APPLICATION EXPERIENCE · 37 RECORDS", styles["h2"]),
    ]
    for section in data["experience_sections"]:
        story.append(RLParagraph(html.escape(section["title"]), styles["h2"]))
        for record in section["records"]:
            line = (
                f"<b>{html.escape(record['role'])}</b> | "
                f"{html.escape(record['description'])} | "
                f"<i>{html.escape(record['organization'])}</i> | "
                f"{html.escape(record['dates'])}"
            )
            story.append(RLParagraph(line, styles["body"]))

    story.extend(
        [
            RLParagraph("EDUCATION", styles["h2"]),
            *[
                RLParagraph("• " + html.escape(item), styles["compact"])
                for item in data["education"]
            ],
            RLParagraph("CREDENTIALS & PROFESSIONAL LEARNING", styles["h2"]),
            *[
                RLParagraph("• " + html.escape(item), styles["compact"])
                for item in data["credentials"] + data["professional_learning"]
            ],
            RLParagraph("METHODS, TOOLS & LANGUAGES", styles["h2"]),
            RLParagraph(
                html.escape(
                    " | ".join(data["methods_tools"] + data["languages"])
                ),
                styles["body"],
            ),
            RLParagraph("COURSES & PROGRAMS TAUGHT", styles["h2"]),
            RLParagraph(
                html.escape(
                    " | ".join(
                        f"{code} {title}"
                        for code, title in data["courses"]["ontario"]
                    )
                ),
                styles["body"],
            ),
            RLParagraph(
                html.escape(" | ".join(data["courses"]["international"])),
                styles["body"],
            ),
            RLParagraph("HISTORICAL CAREER & PROJECT ARCHIVE", styles["subtitle"]),
            RLParagraph(
                "The archive preserves the broader record separately from the "
                "37-row application CV.",
                styles["body"],
            ),
        ]
    )
    for record in historical_records(data):
        title = html.escape(record.get("title") or "Untitled record")
        date = html.escape(record.get("date") or "")
        note = html.escape(record.get("note") or "")
        description = html.escape(record.get("description") or "")
        block = [
            RLParagraph(title, styles["h3"]),
            RLParagraph(
                " | ".join(value for value in [date, note] if value),
                styles["meta"],
            ),
        ]
        if description:
            block.append(RLParagraph(description, styles["body"]))
        story.append(KeepTogether(block))
        story.append(Spacer(1, 2))

    document.build(story, onFirstPage=page_frame, onLaterPages=page_frame)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_pages(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def package_outputs(paths: list[Path], output: Path) -> None:
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(set(paths), key=lambda item: item.name):
            info = zipfile.ZipInfo(path.name, (2026, 7, 28, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, path.read_bytes())


def build_all() -> dict:
    register_pdf_fonts()
    data = load_data()
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    proton_docx = DOWNLOADS / (
        "saul-karim-nassau-ultimate-school-cv-2026-protonmail.docx"
    )
    gmail_docx = DOWNLOADS / "saul-karim-nassau-ultimate-school-cv-2026-gmail.docx"

    chosen_leading = None
    proton_pdf = proton_docx.with_suffix(".pdf")
    for row_leading in (12.5, 12.3, 12.1, 11.9):
        build_application_docx(
            data,
            proton_docx,
            data["contact"]["public_email"],
            row_leading=row_leading,
        )
        convert_docx_to_pdf(proton_docx)
        if pdf_pages(proton_pdf) == 1:
            chosen_leading = row_leading
            break
    if chosen_leading is None:
        raise RuntimeError("The 37-row application CV did not fit on one page.")

    clone_docx_email_only(
        proton_docx,
        gmail_docx,
        data["contact"]["public_email"],
        data["contact"]["alternate_email"],
    )
    gmail_pdf = convert_docx_to_pdf(gmail_docx)
    if pdf_pages(gmail_pdf) != 1:
        raise RuntimeError("The Gmail application CV did not fit on one page.")

    generated: list[Path] = [proton_docx, proton_pdf, gmail_docx, gmail_pdf]
    modular_outputs = []
    for module in data["focus_modules"]:
        slug = module["id"]
        designed = DOWNLOADS / f"saul-karim-nassau-{slug}-cv.pdf"
        ats = DOWNLOADS / f"saul-karim-nassau-{slug}-cv-ats.pdf"
        text_path = DOWNLOADS / f"saul-karim-nassau-{slug}-cv.txt"
        modular_pdf(data, module, designed, ats=False)
        modular_pdf(data, module, ats, ats=True)
        text_path.write_text(modular_text(data, module), encoding="utf-8")
        generated.extend([designed, ats, text_path])
        modular_outputs.extend(
            [
                {
                    "focus": slug,
                    "type": "designed",
                    "path": designed.relative_to(ROOT).as_posix(),
                    "pages": pdf_pages(designed),
                    "bytes": designed.stat().st_size,
                    "sha256": sha256(designed),
                },
                {
                    "focus": slug,
                    "type": "ats",
                    "path": ats.relative_to(ROOT).as_posix(),
                    "pages": pdf_pages(ats),
                    "bytes": ats.stat().st_size,
                    "sha256": sha256(ats),
                },
                {
                    "focus": slug,
                    "type": "text",
                    "path": text_path.relative_to(ROOT).as_posix(),
                    "bytes": text_path.stat().st_size,
                    "sha256": sha256(text_path),
                },
            ]
        )

    everything = DOWNLOADS / "saul-karim-nassau-complete-career-archive-cv.pdf"
    everything_pdf(data, everything)
    generated.append(everything)

    all_zip = DOWNLOADS / "saul-karim-nassau-all-cv-outputs.zip"
    package_outputs(generated, all_zip)

    manifest = {
        "release": data["release"],
        "status": "active",
        "source_of_truth": DATA_PATH.relative_to(ROOT).as_posix(),
        "template": TEMPLATE_PATH.relative_to(ROOT).as_posix(),
        "application_experience_rows": len(all_application_records(data)),
        "application_sections": section_counts(data),
        "application_row_leading_points": chosen_leading,
        "policy": {
            "one_edit_rebuilds_every_output": True,
            "gmail_and_protonmail_differ_only_by_email": True,
            "role_focused_pdfs_are_distinct_outputs": True,
            "everything_pdf_contains_application_and_historical_records": True,
        },
        "application_outputs": [
            {
                "edition": "protonmail",
                "format": "docx",
                "email": data["contact"]["public_email"],
                "path": proton_docx.relative_to(ROOT).as_posix(),
                "pages": 1,
                "bytes": proton_docx.stat().st_size,
                "sha256": sha256(proton_docx),
            },
            {
                "edition": "protonmail",
                "format": "pdf",
                "email": data["contact"]["public_email"],
                "path": proton_pdf.relative_to(ROOT).as_posix(),
                "pages": pdf_pages(proton_pdf),
                "bytes": proton_pdf.stat().st_size,
                "sha256": sha256(proton_pdf),
            },
            {
                "edition": "gmail",
                "format": "docx",
                "email": data["contact"]["alternate_email"],
                "path": gmail_docx.relative_to(ROOT).as_posix(),
                "pages": 1,
                "bytes": gmail_docx.stat().st_size,
                "sha256": sha256(gmail_docx),
            },
            {
                "edition": "gmail",
                "format": "pdf",
                "email": data["contact"]["alternate_email"],
                "path": gmail_pdf.relative_to(ROOT).as_posix(),
                "pages": pdf_pages(gmail_pdf),
                "bytes": gmail_pdf.stat().st_size,
                "sha256": sha256(gmail_pdf),
            },
        ],
        "modular_outputs": modular_outputs,
        "everything_output": {
            "path": everything.relative_to(ROOT).as_posix(),
            "pages": pdf_pages(everything),
            "bytes": everything.stat().st_size,
            "sha256": sha256(everything),
        },
        "bundle": {
            "path": all_zip.relative_to(ROOT).as_posix(),
            "bytes": all_zip.stat().st_size,
            "sha256": sha256(all_zip),
        },
    }
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


if __name__ == "__main__":
    result = build_all()
    print(
        "Built all Saul CV outputs from one source: "
        f"{result['application_experience_rows']} application experiences, "
        f"{len(result['modular_outputs'])} modular outputs."
    )
